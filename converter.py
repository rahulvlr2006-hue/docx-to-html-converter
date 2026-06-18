from docx import Document


def convert_docx_to_html(path):

    doc = Document(path)

    html = ""
    inside_list = False

    # Paragraphs
    for para in doc.paragraphs:

        if not para.text.strip():
            continue

        if para.text.startswith("•"):

            if not inside_list:
                html += "<ul>\n"
                inside_list = True

            item = para.text.replace("•", "").strip()
            html += f"    <li>{item}</li>\n"

            continue

        else:
            if inside_list:
                html += "</ul>\n"
                inside_list = False

        style = para.style.name
        paragraph_html = ""

        for run in para.runs:

            text = run.text

            if run.bold:
                text = f"<strong>{text}</strong>"

            if run.italic:
                text = f"<em>{text}</em>"

            if run.underline:
                text = f"<u>{text}</u>"

            paragraph_html += text

        if style == "Heading 1":
            html += f"<h1>{paragraph_html}</h1>\n"

        elif style == "Heading 2":
            html += f"<h2>{paragraph_html}</h2>\n"

        elif style == "Heading 3":
            html += f"<h3>{paragraph_html}</h3>\n"

        else:
            html += f"<p>{paragraph_html}</p>\n"

    if inside_list:
        html += "</ul>\n"

    # Tables
    for table in doc.tables:

        html += "<table border='1'>\n"

        for row_index, row in enumerate(table.rows):

            html += "    <tr>\n"

            for cell in row.cells:

                if row_index == 0:
                    html += f"        <th>{cell.text}</th>\n"
                else:
                    html += f"        <td>{cell.text}</td>\n"

            html += "    </tr>\n"

        html += "</table>\n"

    return html